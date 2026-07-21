"""Reporting queries + tabular exports (CSV / XLSX / PDF)."""
from __future__ import annotations

import csv
import io
from datetime import date
from decimal import Decimal
from typing import Any

from sqlalchemy import Select, func, select
from sqlalchemy.orm import Session

from app.models.enums import ValidationStatus
from app.models.fakturownia import (
    FakturowniaInvoice,
    FakturowniaInvoicePosition,
    FakturowniaProduct,
    FakturowniaProductStock,
    Warehouse,
)
from app.models.product import Product
from app.models.stock import LocalStockBalance, LocalStockLedgerEntry
from app.models.validation import ValidationIssue

OPEN_ISSUE_STATUSES = [
    str(ValidationStatus.WARNING),
    str(ValidationStatus.ERROR),
    str(ValidationStatus.NEEDS_MAPPING),
    str(ValidationStatus.NEEDS_OPENING_BALANCE),
]


# --------------------------------------------------------------------- #
# Report queries — each returns (headers, rows)
# --------------------------------------------------------------------- #
def sales_report(
    db: Session, date_from: date | None = None, date_to: date | None = None
) -> tuple[list[str], list[list[Any]]]:
    query: Select = (
        select(
            Product.id,
            Product.name,
            Product.code,
            Product.unit,
            Product.kind,
            func.sum(FakturowniaInvoicePosition.quantity).label("total_quantity"),
            func.sum(FakturowniaInvoicePosition.total_price_net).label("total_net"),
            func.count(func.distinct(FakturowniaInvoice.id)).label("invoice_count"),
        )
        .join(
            FakturowniaInvoicePosition,
            FakturowniaInvoicePosition.mapped_product_id == Product.id,
        )
        .join(
            FakturowniaInvoice,
            FakturowniaInvoice.id == FakturowniaInvoicePosition.invoice_id,
        )
        .where(FakturowniaInvoice.is_cancelled.is_(False))
        .group_by(Product.id, Product.name, Product.code, Product.unit, Product.kind)
        .order_by(func.sum(FakturowniaInvoicePosition.total_price_net).desc())
    )
    if date_from:
        query = query.where(FakturowniaInvoice.sell_date >= date_from)
    if date_to:
        query = query.where(FakturowniaInvoice.sell_date <= date_to)
    headers = [
        "ID produktu", "Nazwa", "Kod", "Jednostka", "Typ",
        "Ilość sprzedana", "Wartość netto", "Liczba faktur",
    ]
    rows = [list(row) for row in db.execute(query).all()]
    return headers, rows


def unmapped_products_report(db: Session) -> tuple[list[str], list[list[Any]]]:
    query = (
        select(
            FakturowniaInvoicePosition.name,
            FakturowniaInvoicePosition.code,
            FakturowniaInvoicePosition.quantity_unit,
            func.count(FakturowniaInvoicePosition.id).label("occurrences"),
            func.sum(FakturowniaInvoicePosition.quantity).label("total_quantity"),
            func.sum(FakturowniaInvoicePosition.total_price_net).label("total_net"),
        )
        .where(FakturowniaInvoicePosition.mapped_product_id.is_(None))
        .group_by(
            FakturowniaInvoicePosition.name,
            FakturowniaInvoicePosition.code,
            FakturowniaInvoicePosition.quantity_unit,
        )
        .order_by(func.count(FakturowniaInvoicePosition.id).desc())
    )
    headers = ["Nazwa pozycji", "Kod", "Jednostka", "Wystąpienia", "Ilość łącznie", "Wartość netto"]
    rows = [list(row) for row in db.execute(query).all()]
    return headers, rows


def negative_stock_report(db: Session) -> tuple[list[str], list[list[Any]]]:
    query = (
        select(
            Product.name,
            Product.code,
            Warehouse.name,
            LocalStockBalance.quantity,
            LocalStockBalance.last_movement_at,
            LocalStockBalance.last_sale_at,
        )
        .join(Product, Product.id == LocalStockBalance.product_id)
        .join(Warehouse, Warehouse.id == LocalStockBalance.warehouse_id)
        .where(LocalStockBalance.quantity < 0)
        .order_by(LocalStockBalance.quantity)
    )
    headers = ["Produkt", "Kod", "Magazyn", "Stan lokalny", "Ostatni ruch", "Ostatnia sprzedaż"]
    rows = [list(row) for row in db.execute(query).all()]
    return headers, rows


def stock_discrepancies_report(db: Session) -> tuple[list[str], list[list[Any]]]:
    headers = ["Produkt", "Kod", "Magazyn", "Stan lokalny", "Stan Fakturownia", "Różnica"]
    rows: list[list[Any]] = []
    remote_stock: dict[tuple[int, int], Decimal] = {}
    for row in db.execute(select(FakturowniaProductStock)).scalars():
        if row.quantity is not None:
            remote_stock[(row.product_fakturownia_id, row.warehouse_fakturownia_id)] = row.quantity
    remote_global = {
        fp.fakturownia_id: fp.quantity
        for fp in db.execute(select(FakturowniaProduct)).scalars()
        if fp.quantity is not None
    }
    warehouses = {w.id: w for w in db.execute(select(Warehouse)).scalars()}
    products = {prod.id: prod for prod in db.execute(select(Product)).scalars()}

    for balance in db.execute(select(LocalStockBalance)).scalars():
        product = products.get(balance.product_id)
        warehouse = warehouses.get(balance.warehouse_id)
        if product is None or product.fakturownia_product_id is None:
            continue
        remote_qty = None
        if warehouse is not None and warehouse.fakturownia_id is not None:
            remote_qty = remote_stock.get(
                (product.fakturownia_product_id, warehouse.fakturownia_id)
            )
        if remote_qty is None:
            remote_qty = remote_global.get(product.fakturownia_product_id)
        if remote_qty is None:
            continue
        difference = balance.quantity - remote_qty
        if difference != 0:
            rows.append([
                product.name, product.code,
                warehouse.name if warehouse else balance.warehouse_id,
                balance.quantity, remote_qty, difference,
            ])
    rows.sort(key=lambda r: abs(r[5] or 0), reverse=True)
    return headers, rows


def problem_invoices_report(db: Session) -> tuple[list[str], list[list[Any]]]:
    query = (
        select(
            FakturowniaInvoice.number,
            FakturowniaInvoice.issue_date,
            FakturowniaInvoice.buyer_name,
            FakturowniaInvoice.kind,
            FakturowniaInvoice.total_gross,
            func.count(ValidationIssue.id).label("open_issues"),
        )
        .join(ValidationIssue, ValidationIssue.invoice_id == FakturowniaInvoice.id)
        .where(ValidationIssue.status.in_(OPEN_ISSUE_STATUSES))
        .group_by(
            FakturowniaInvoice.id, FakturowniaInvoice.number,
            FakturowniaInvoice.issue_date, FakturowniaInvoice.buyer_name,
            FakturowniaInvoice.kind, FakturowniaInvoice.total_gross,
        )
        .order_by(func.count(ValidationIssue.id).desc())
    )
    headers = ["Numer", "Data wystawienia", "Nabywca", "Typ", "Wartość brutto", "Otwarte problemy"]
    rows = [list(row) for row in db.execute(query).all()]
    return headers, rows


def top_error_products_report(db: Session) -> tuple[list[str], list[list[Any]]]:
    query = (
        select(
            Product.name,
            Product.code,
            ValidationIssue.issue_type,
            func.count(ValidationIssue.id).label("issue_count"),
        )
        .join(Product, Product.id == ValidationIssue.product_id)
        .group_by(Product.id, Product.name, Product.code, ValidationIssue.issue_type)
        .order_by(func.count(ValidationIssue.id).desc())
        .limit(100)
    )
    headers = ["Produkt", "Kod", "Typ problemu", "Liczba problemów"]
    rows = [list(row) for row in db.execute(query).all()]
    return headers, rows


def stock_history_report(
    db: Session, product_id: int | None = None, warehouse_id: int | None = None
) -> tuple[list[str], list[list[Any]]]:
    query = (
        select(
            LocalStockLedgerEntry.occurred_at,
            Product.name,
            Warehouse.name,
            LocalStockLedgerEntry.movement_type,
            LocalStockLedgerEntry.quantity_change,
            LocalStockLedgerEntry.balance_before,
            LocalStockLedgerEntry.balance_after,
            LocalStockLedgerEntry.description,
        )
        .join(Product, Product.id == LocalStockLedgerEntry.product_id)
        .join(Warehouse, Warehouse.id == LocalStockLedgerEntry.warehouse_id)
        .order_by(LocalStockLedgerEntry.sequence)
    )
    if product_id:
        query = query.where(LocalStockLedgerEntry.product_id == product_id)
    if warehouse_id:
        query = query.where(LocalStockLedgerEntry.warehouse_id == warehouse_id)
    headers = [
        "Data", "Produkt", "Magazyn", "Typ ruchu", "Zmiana",
        "Stan przed", "Stan po", "Opis",
    ]
    rows = [list(row) for row in db.execute(query).all()]
    return headers, rows


# --------------------------------------------------------------------- #
# Export helpers
# --------------------------------------------------------------------- #
def _stringify(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    return value


def export_csv(headers: list[str], rows: list[list[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_stringify(v) for v in row])
    return buffer.getvalue().encode("utf-8-sig")


def export_xlsx(headers: list[str], rows: list[list[Any]], title: str = "Raport") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            sheet.cell(row=row_idx, column=col_idx, value=_stringify(value))
    for col_idx, header in enumerate(headers, start=1):
        width = max(
            len(str(header)),
            *(len(str(_stringify(r[col_idx - 1]))) for r in rows[:200]) if rows else (10,),
        )
        sheet.column_dimensions[sheet.cell(row=1, column=col_idx).column_letter].width = min(
            width + 2, 60
        )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_pdf(headers: list[str], rows: list[list[Any]], title: str = "Raport") -> bytes:
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, _latin(title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 7)
    page_width = pdf.w - pdf.l_margin - pdf.r_margin
    col_width = page_width / max(len(headers), 1)
    for header in headers:
        pdf.cell(col_width, 6, _latin(str(header))[:30], border=1)
    pdf.ln()
    pdf.set_font("Helvetica", "", 7)
    for row in rows[:2000]:
        for value in row:
            pdf.cell(col_width, 5, _latin(str(_stringify(value)))[:30], border=1)
        pdf.ln()
    return bytes(pdf.output())


def _latin(text: str) -> str:
    """fpdf core fonts are latin-1 only; replace unsupported characters."""
    return text.encode("latin-1", errors="replace").decode("latin-1")


REPORTS = {
    "sales": sales_report,
    "unmapped-products": unmapped_products_report,
    "negative-stock": negative_stock_report,
    "stock-discrepancies": stock_discrepancies_report,
    "problem-invoices": problem_invoices_report,
    "top-error-products": top_error_products_report,
    "stock-history": stock_history_report,
}
